"""
Reportes de actividad por usuario (Pagos, Compras y Ventas registrados), por rango de fechas,
en Excel y PDF. Sirve tanto para un usuario como para el reporte general de todos.

Flujo:
  1. app.py trae los registros del rango (db.listar_actividad_rango).
  2. agrupar_actividad_por_usuario() los separa por "encargado" y los convierte en filas listas
     para mostrar (columnas en español, "Monto" numérico).
  3. construir_reporte_usuario() suma los totales; generar_excel_reporte()/generar_pdf_reporte()
     solo maquetan.

Alcance: solo se pueden atribuir a un usuario los movimientos que guardan "encargado": Pagos,
Compras y Ventas. Los abonos de alquiler/electricidad/agua no registran quién los cobró.
Totales: Egresos = pagos + compras; Ingresos = ventas (no se mezclan en una sola cifra).
"""
import io
from datetime import datetime
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
                                CondPageBreak, HRFlowable, Image as RLImage)

import recibo

# (clave_interna, título, tipo)   tipo: "egreso" | "ingreso"
CATEGORIAS = [
    ("pagos_generales", "Pagos realizados", "egreso"),
    ("compras", "Compras realizadas", "egreso"),
    ("ventas", "Ventas realizadas", "ingreso"),
]
SIN_ENCARGADO = "(Sin encargado)"

_COLOR_HEADER = "2E4057"   # sin "#" para openpyxl; con "#" para reportlab
_FMT_EXCEL = '"Bs" #,##0.00'
_PELIGRO_FORMULA = ("=", "+", "-", "@")


# ============================================================
# Datos: de registros de la base a filas de reporte
# ============================================================

def _fmt_bs(v):
    try:
        return f"Bs {float(v):,.2f}"
    except (TypeError, ValueError):
        return "Bs 0.00"


def _fmt_fecha(valor):
    texto = str(valor or "")[:10]
    try:
        return datetime.strptime(texto, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return texto or "—"


def _txt(valor):
    return str(valor).strip() if valor not in (None, "") else ""


def _estado_txt(valor):
    """Texto del estado para el reporte: '—' si no aplica (o si el registro es anterior al campo)."""
    texto = _txt(valor)
    return texto if texto and texto != "No aplica" else "—"


def fila_pago_general(r):
    return {"Fecha": _fmt_fecha(r.get("fecha_pago")), "Categoría": _txt(r.get("categoria")),
            "Descripción": _txt(r.get("descripcion")), "Beneficiario": _txt(r.get("beneficiario")),
            "Método de pago": _txt(r.get("metodo_pago")), "Estado": _estado_txt(r.get("estado_devolucion")),
            "Monto": float(r.get("monto") or 0)}


def fila_compra(r):
    return {"Fecha": _fmt_fecha(r.get("fecha_compra")), "Categoría": _txt(r.get("categoria")),
            "Descripción": _txt(r.get("descripcion")), "Proveedor": _txt(r.get("proveedor")),
            "Comprobante N°": _txt(r.get("numero_comprobante")),
            "Método de pago": _txt(r.get("metodo_pago")), "Estado": _estado_txt(r.get("estado_devolucion")),
            "Monto": float(r.get("monto_total") or 0)}


def fila_venta(r):
    return {"Fecha": _fmt_fecha(r.get("fecha_venta")), "Concepto": _txt(r.get("concepto")),
            "Descripción": _txt(r.get("descripcion")), "Comprador": _txt(r.get("comprador")),
            "Recibo N°": _txt(r.get("recibo_emitido")),
            "Forma de cobro": _txt(r.get("forma_cobro")), "Estado": _estado_txt(r.get("estado_entrega")),
            "Monto": float(r.get("monto") or 0)}


def agrupar_actividad_por_usuario(pagos_generales, compras, ventas, usuarios_extra=()):
    """Devuelve {nombre_usuario: {clave_categoria: [filas]}}, ordenado por nombre.

    Se agrupa por el texto de "encargado" (así se guardó en cada registro). 'usuarios_extra' son
    nombres que deben aparecer aunque no tengan actividad (ej. usuarios activos). Los registros
    sin encargado van al grupo SIN_ENCARGADO, para que los totales generales cuadren con todo
    lo registrado."""
    grupos = {}

    def _grupo(nombre):
        nombre = _txt(nombre) or SIN_ENCARGADO
        return grupos.setdefault(nombre, {clave: [] for clave, _t, _tipo in CATEGORIAS})

    for nombre in usuarios_extra:
        if _txt(nombre):
            _grupo(nombre)
    for registros, clave, mapear in ((pagos_generales, "pagos_generales", fila_pago_general),
                                     (compras, "compras", fila_compra),
                                     (ventas, "ventas", fila_venta)):
        for r in registros or []:
            _grupo(r.get("encargado"))[clave].append(mapear(r))
    return dict(sorted(grupos.items(), key=lambda kv: (kv[0] == SIN_ENCARGADO, kv[0].casefold())))


def construir_reporte_usuario(nombre_usuario, datos_por_categoria):
    """Cada fila debe tener "Monto" (número).
    -> {"usuario", "categorias": {clave: {"titulo","tipo","filas","total"}},
        "total_egresos", "total_ingresos"}"""
    categorias, egresos, ingresos = {}, 0.0, 0.0
    for clave, titulo, tipo in CATEGORIAS:
        filas = datos_por_categoria.get(clave, [])
        total = round(sum(float(f.get("Monto") or 0) for f in filas), 2)
        categorias[clave] = {"titulo": titulo, "tipo": tipo, "filas": filas, "total": total}
        if tipo == "egreso":
            egresos += total
        else:
            ingresos += total
    return {"usuario": nombre_usuario, "categorias": categorias,
            "total_egresos": round(egresos, 2), "total_ingresos": round(ingresos, 2)}


def _pdf_texto(valor):
    """Texto seguro para Paragraph de reportlab: escapa <, > y &, y cambia por '?' lo que la
    fuente estándar (Helvetica) no puede dibujar (emojis, etc.) para no mostrar cuadros negros."""
    return escape(str(valor).encode("cp1252", "replace").decode("cp1252"))


def _nombre_hoja_valido(nombre, usados):
    """Nombres de hoja de Excel: sin \\ / * ? : [ ], únicos y de máx 31 caracteres."""
    limpio = "".join(c for c in nombre if c not in r"\/*?:[]") or "Usuario"
    limpio = limpio[:28]
    candidato, i = limpio, 2
    while candidato.casefold() in {u.casefold() for u in usados}:
        candidato = f"{limpio[:25]} {i}"
        i += 1
    usados.add(candidato)
    return candidato


# ============================================================
# EXCEL
# ============================================================

def _poner(ws, fila, col, valor, fuente=None, formato=None):
    """Escribe una celda. Un texto que empiece con = + - @ se guarda SIEMPRE como texto: si no,
    openpyxl lo convertiría en fórmula y una descripción como '=1+1' (o algo malicioso escrito
    en un campo libre) se ejecutaría al abrir el archivo."""
    c = ws.cell(fila, col, valor)
    if isinstance(valor, str) and valor.startswith(_PELIGRO_FORMULA):
        c.data_type = "s"
    if fuente:
        c.font = fuente
    if formato:
        c.number_format = formato
    return c


def generar_excel_reporte(reportes, fecha_desde, fecha_hasta, emitido_por=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    f_titulo = Font(name="Calibri", bold=True, size=14)
    f_header = Font(name="Calibri", bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor=_COLOR_HEADER)
    f_normal = Font(name="Calibri", size=10)
    f_bold = Font(name="Calibri", bold=True, size=10)
    f_total = Font(name="Calibri", bold=True, size=12)

    # ---------- Resumen ----------
    ws = wb.create_sheet("Resumen")
    _poner(ws, 1, 1, "Reporte de actividad por usuario", f_titulo)
    _poner(ws, 2, 1, f"Periodo: {fecha_desde} al {fecha_hasta}", f_normal)
    if emitido_por:
        _poner(ws, 3, 1, f"Emitido por: {emitido_por}", f_normal)
    encabezados = ["Usuario"] + [t for _c, t, _tp in CATEGORIAS] + ["Egresos (pagos + compras)", "Ingresos (ventas)"]
    fila = 5
    for col, texto in enumerate(encabezados, start=1):
        c = _poner(ws, fila, col, texto, f_header)
        c.fill = relleno
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 32
    fila += 1
    for r in reportes:
        _poner(ws, fila, 1, r["usuario"], f_normal)
        for col, (clave, _t, _tp) in enumerate(CATEGORIAS, start=2):
            _poner(ws, fila, col, r["categorias"][clave]["total"], f_normal, _FMT_EXCEL)
        _poner(ws, fila, len(encabezados) - 1, r["total_egresos"], f_bold, _FMT_EXCEL)
        _poner(ws, fila, len(encabezados), r["total_ingresos"], f_bold, _FMT_EXCEL)
        fila += 1
    if len(reportes) > 1:
        fila += 1
        _poner(ws, fila, 1, "TOTAL GENERAL", f_total)
        for col, (clave, _t, _tp) in enumerate(CATEGORIAS, start=2):
            _poner(ws, fila, col, round(sum(r["categorias"][clave]["total"] for r in reportes), 2), f_bold, _FMT_EXCEL)
        _poner(ws, fila, len(encabezados) - 1, round(sum(r["total_egresos"] for r in reportes), 2), f_total, _FMT_EXCEL)
        _poner(ws, fila, len(encabezados), round(sum(r["total_ingresos"] for r in reportes), 2), f_total, _FMT_EXCEL)
    ws.column_dimensions["A"].width = max(26, min(45, max([len(r["usuario"]) for r in reportes] + [0]) + 2))
    for col in range(2, len(encabezados) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    # ---------- Una hoja por usuario ----------
    usados = {"Resumen"}
    for r in reportes:
        ws = wb.create_sheet(_nombre_hoja_valido(r["usuario"] or "Usuario", usados))
        _poner(ws, 1, 1, f'Detalle de {r["usuario"]}', f_titulo)
        _poner(ws, 2, 1, f"Periodo: {fecha_desde} al {fecha_hasta}", f_normal)
        fila = 4
        anchos = {}
        for clave, titulo, _tipo in CATEGORIAS:
            datos = r["categorias"][clave]
            _poner(ws, fila, 1, titulo, Font(name="Calibri", bold=True, size=12))
            fila += 1
            if not datos["filas"]:
                _poner(ws, fila, 1, "Sin registros en este periodo.", f_normal)
                fila += 2
                continue
            columnas = list(datos["filas"][0].keys())
            for col, nombre_col in enumerate(columnas, start=1):
                c = _poner(ws, fila, col, nombre_col, f_header)
                c.fill = relleno
                anchos[col] = max(anchos.get(col, 0), len(nombre_col))
            fila += 1
            for fd in datos["filas"]:
                for col, nombre_col in enumerate(columnas, start=1):
                    valor = fd[nombre_col]
                    c = _poner(ws, fila, col, valor, f_normal, _FMT_EXCEL if nombre_col == "Monto" else None)
                    if nombre_col == "Descripción":
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                    anchos[col] = max(anchos.get(col, 0), len(str(valor)) if nombre_col != "Monto" else 14)
                fila += 1
            _poner(ws, fila, 1, "Subtotal:", f_bold)
            _poner(ws, fila, 2, datos["total"], f_bold, _FMT_EXCEL)
            fila += 2
        _poner(ws, fila, 1, "TOTAL EGRESOS (pagos + compras):", f_total)
        _poner(ws, fila, 4, r["total_egresos"], f_total, _FMT_EXCEL)
        fila += 1
        _poner(ws, fila, 1, "TOTAL INGRESOS (ventas):", f_total)
        _poner(ws, fila, 4, r["total_ingresos"], f_total, _FMT_EXCEL)
        for col in range(1, 9):
            ws.column_dimensions[get_column_letter(col)].width = max(14, min(45, anchos.get(col, 14) + 2))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ============================================================
# PDF (horizontal; mismo estilo visual que el recibo)
# ============================================================

def generar_pdf_reporte(reportes, fecha_desde, fecha_hasta, emitido_por=None):
    buffer = io.BytesIO()
    pagina = landscape(letter)
    doc = SimpleDocTemplate(buffer, pagesize=pagina, topMargin=12 * mm, bottomMargin=12 * mm,
                             leftMargin=12 * mm, rightMargin=12 * mm, title="Reporte de actividad por usuario")
    ancho_util = pagina[0] - 24 * mm
    base = getSampleStyleSheet()["Normal"]
    st_titulo = ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=18)
    st_sub = ParagraphStyle("s", parent=base, fontSize=9, leading=12, textColor=colors.Color(.33, .33, .33))
    st_chico = ParagraphStyle("c", parent=base, fontSize=8, leading=10)
    st_cel = ParagraphStyle("cel", parent=base, fontSize=7, leading=8.5)
    st_cel_der = ParagraphStyle("celd", parent=st_cel, alignment=2)
    st_h = ParagraphStyle("h", parent=st_cel, fontName="Helvetica-Bold", textColor=colors.white)
    st_h_der = ParagraphStyle("hd", parent=st_h, alignment=2)
    st_usuario = ParagraphStyle("u", parent=base, fontName="Helvetica-Bold", fontSize=12, leading=15, spaceBefore=6)
    st_cat = ParagraphStyle("cat", parent=base, fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=6, spaceAfter=2)
    st_vacio = ParagraphStyle("v", parent=base, fontName="Helvetica-Oblique", fontSize=8, textColor=colors.Color(.33, .33, .33))
    story = []

    # ---------- Encabezado: logo + título, como el recibo ----------
    logo = ""
    try:
        import os
        if os.path.exists(recibo.LOGO_PATH):
            logo = RLImage(io.BytesIO(recibo._logo_liviano_bytes()), width=16 * mm, height=16 * mm)
    except Exception:
        logo = ""
    cab = Table([[logo, [Paragraph(recibo.NOMBRE_RESIDENCIAL, st_titulo),
                         Paragraph("Reporte de actividad por usuario", st_sub)]]],
                colWidths=[20 * mm, ancho_util - 20 * mm])
    cab.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    emitido = datetime.now(recibo._ZONA_BOLIVIA).strftime("%d/%m/%Y %H:%M")
    story += [cab, Spacer(1, 3), HRFlowable(width="100%", thickness=0.6, color=colors.black), Spacer(1, 4),
              Paragraph(f"<b>Periodo:</b> {_pdf_texto(str(fecha_desde))} al {_pdf_texto(str(fecha_hasta))}"
                        f" &nbsp;&nbsp;|&nbsp;&nbsp; <b>Emitido:</b> {emitido}"
                        + (f" &nbsp;&nbsp;|&nbsp;&nbsp; <b>Por:</b> {_pdf_texto(str(emitido_por))}" if emitido_por else ""),
                        st_chico),
              Spacer(1, 8)]

    # ---------- Resumen ----------
    encabezado = ["Usuario"] + [t for _c, t, _tp in CATEGORIAS] + ["Egresos (pagos + compras)", "Ingresos (ventas)"]
    filas_res = [[Paragraph(h, st_h if i == 0 else st_h_der) for i, h in enumerate(encabezado)]]
    for r in reportes:
        filas_res.append([Paragraph(_pdf_texto(r["usuario"]), st_cel)] +
                         [Paragraph(_fmt_bs(r["categorias"][c]["total"]), st_cel_der) for c, _t, _tp in CATEGORIAS] +
                         [Paragraph(f'<b>{_fmt_bs(r["total_egresos"])}</b>', st_cel_der),
                          Paragraph(f'<b>{_fmt_bs(r["total_ingresos"])}</b>', st_cel_der)])
    if len(reportes) > 1:
        filas_res.append([Paragraph("<b>TOTAL GENERAL</b>", st_cel)] +
                         [Paragraph(f'<b>{_fmt_bs(sum(r["categorias"][c]["total"] for r in reportes))}</b>', st_cel_der)
                          for c, _t, _tp in CATEGORIAS] +
                         [Paragraph(f'<b>{_fmt_bs(sum(r["total_egresos"] for r in reportes))}</b>', st_cel_der),
                          Paragraph(f'<b>{_fmt_bs(sum(r["total_ingresos"] for r in reportes))}</b>', st_cel_der)])
    t = Table(filas_res, repeatRows=1, hAlign="LEFT",
              colWidths=[ancho_util * w for w in (0.28, 0.12, 0.12, 0.12, 0.18, 0.18)])
    estilo = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + _COLOR_HEADER)),
              ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")),
              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
              ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
    if len(reportes) > 1:
        estilo += [("LINEABOVE", (0, -1), (-1, -1), 1, colors.black), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f0f0f0"))]
    t.setStyle(TableStyle(estilo))
    story += [t, Spacer(1, 10)]

    # ---------- Detalle por usuario ----------
    peso = {"Descripción": 3.2, "Fecha": 1.1, "Monto": 1.2}
    for r in reportes:
        story.append(CondPageBreak(45 * mm))
        story.append(Paragraph(f'Detalle: {_pdf_texto(r["usuario"])}', st_usuario))
        for clave, titulo, _tipo in CATEGORIAS:
            datos = r["categorias"][clave]
            story.append(CondPageBreak(30 * mm))
            story.append(Paragraph(titulo, st_cat))
            if not datos["filas"]:
                story.append(Paragraph("Sin registros en este periodo.", st_vacio))
                continue
            columnas = list(datos["filas"][0].keys())
            pesos = [peso.get(c, 1.6) for c in columnas]
            anchos = [ancho_util * p / sum(pesos) for p in pesos]
            tabla = [[Paragraph(_pdf_texto(c), st_h_der if c == "Monto" else st_h) for c in columnas]]
            for f in datos["filas"]:
                tabla.append([Paragraph(_fmt_bs(f[c]) if c == "Monto" else _pdf_texto(str(f[c])),
                                        st_cel_der if c == "Monto" else st_cel) for c in columnas])
            tt = Table(tabla, repeatRows=1, hAlign="LEFT", colWidths=anchos)
            tt.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + _COLOR_HEADER)),
                ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]))
            story.append(tt)
            story.append(Paragraph(f'<para align="right"><b>Subtotal: {_fmt_bs(datos["total"])}</b></para>', st_chico))
        story.append(Spacer(1, 4))
        story.append(Paragraph(
            f'<para align="right"><b>Egresos (pagos + compras): {_fmt_bs(r["total_egresos"])} &nbsp;&nbsp;|&nbsp;&nbsp; '
            f'Ingresos (ventas): {_fmt_bs(r["total_ingresos"])}</b></para>', st_chico))
        story.append(HRFlowable(width="100%", thickness=0.4, color=colors.HexColor("#cccccc")))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
