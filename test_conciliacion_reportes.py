"""Pruebas de conciliacion_reportes.py  ->  python test_conciliacion_reportes.py   (necesita test_conciliacion.py y recibo.py al lado)"""
import io
import openpyxl
import conciliacion as co
import conciliacion_reportes as rep
import extracto as x
import test_conciliacion as tc
import test_extracto as te

r = tc.r
ext = x.leer_extracto(te.libro_xlsx(), "e.xlsx")
est = rep.estado_extracto(r)
assert len(est) == 27 and sum(1 for v in est.values() if v[0] == "Conciliado") == 15 and sum(1 for v in est.values() if v[0] == "Sin registrar en el sistema") == 12
assert est[2][0] == "Conciliado" and "1 a 1" in est[2][1] and "Alta" in est[2][1] and est[9][1].startswith("1 banco = varios registros")
assert est[17][0] == "Sin registrar en el sistema" and "Probablemente falta registrarlo" in est[17][1]

# ---- filas ----
conc = rep.filas_conciliados(r)
assert len(conc) == 14 and all(f["monto_banco"] == f["monto_sistema"] for f in conc)                      # cada coincidencia suma lo mismo de los dos lados (con signo)
assert round(sum(f["monto_banco"] for f in conc), 2) == round(r["resumen"]["conciliado"]["creditos"] - r["resumen"]["conciliado"]["debitos"], 2)
persianas = [f for f in conc if "Persianas" in f["banco"]][0]
assert persianas["monto_banco"] == -5200 and persianas["monto_sistema"] == -5200
assert persianas["union"] == co.VARIOS_BANCO and " + " in persianas["banco"] and persianas["fecha_banco"] == "07/10/2026, 13/10/2026" and persianas["dias"] == 6
bsr = rep.filas_banco_sin_registro(r)
assert len(bsr) == 12 and [f["fecha"] for f in bsr] == sorted([f["fecha"] for f in bsr], key=lambda s: (s[6:], s[3:5], s[:2]))   # en orden de fecha
assert round(sum(f["monto"] for f in bsr), 2) == round(r["resumen"]["sin_registro"]["creditos"] - r["resumen"]["sin_registro"]["debitos"], 2)
ssb = rep.filas_sistema_sin_banco(r)
assert sorted((f["origen"], f["monto"]) for f in ssb) == [("Alquiler", 1300), ("Compra", -800), ("Luz", 270)] and all(f["causa"] for f in ssb)   # los egresos salen en negativo, como en el banco
assert len(rep.filas_efectivo(r)) == 6 and sorted(f["monto"] for f in rep.filas_efectivo(r))[0] == -100
fx = rep.filas_extracto(r)
assert len(fx) == 27 and fx[0]["n"] == 1 and fx[1]["estado"] == "Conciliado" and fx[3]["estado"] == "Sin registrar en el sistema" and fx[1]["saldo"] == 1100.0

# ---- Excel ----
wb = openpyxl.load_workbook(io.BytesIO(rep.generar_excel(r, ext, "Vio")))
assert wb.sheetnames == ["Resumen", "Conciliados", "Banco sin registrar", "Sistema sin banco", "Efectivo", "Extracto"]
ws = wb["Resumen"]
valores = {ws.cell(i, 1).value: ws.cell(i, 2).value for i in range(1, ws.max_row + 1)}
assert valores["Movimientos del extracto"] == 27 and valores["Créditos (entra dinero)"] == 23510.53 and valores["Débitos (sale dinero)"] == 23461.32
assert valores["Saldo inicial"] == 0.12 and valores["Saldo final"] == 49.33 and valores["Control del saldo del extracto"] == "Correcto"
assert valores["Movimientos del banco conciliados"] == 15 and valores["Movimientos sin registrar"] == 12 and valores["Registros sin movimiento en el banco"] == 3
assert valores["Coincidencias con confianza Alta / Media / Revisar"] == "14 / 0 / 0" and valores["Registros en efectivo"] == 6 and "3 día" in valores["Periodo del extracto"]
assert wb["Banco sin registrar"].max_row == 4 + 12 + 1 and wb["Banco sin registrar"].cell(4 + 12 + 1, 1).value == "TOTAL"
assert wb["Extracto"].max_row == 4 + 27 and [wb["Extracto"].cell(5 + i, 10).value for i in (0, 1)] == ["Sin registrar en el sistema", "Conciliado"]
assert wb["Conciliados"].max_row == 4 + 14
# con el control de saldo con problemas
ext_mal = x.leer_extracto(te.libro_xlsx([f for i, f in enumerate(te.FILAS) if i != 13]), "e.xlsx")
wb2 = openpyxl.load_workbook(io.BytesIO(rep.generar_excel(r, ext_mal)))
assert any(str(wb2["Resumen"].cell(i, 2).value).endswith("revisar") for i in range(1, wb2["Resumen"].max_row + 1))
# el texto del banco o del sistema nunca se interpreta como fórmula
raro = [{**te.x.leer_extracto(te.libro_xlsx(), "e.xlsx")["lineas"][0], "detalle": "=1+1", "monto": 10.0, "fecha": tc.D(2026, 10, 1)}]
rraro = co.conciliar(raro, [])
libro = openpyxl.load_workbook(io.BytesIO(rep.generar_excel(rraro)))
assert all(c.data_type != "f" for hoja in libro for fila in hoja.iter_rows() for c in fila)

# ---- sin efectivo / vacío / todo conciliado ----
sin_ef = co.conciliar(tc.L, tc.sistema_base(), incluir_efectivo=True)
assert "Efectivo" not in openpyxl.load_workbook(io.BytesIO(rep.generar_excel(sin_ef, ext))).sheetnames
vacio = co.conciliar(tc.L, [])
openpyxl.load_workbook(io.BytesIO(rep.generar_excel(vacio, ext)))
todo = co.conciliar(tc.L[1:3], [tc.sis("Alquiler", co.INGRESO, 5, 1100, "PP - APARTAMENTO 3", "A"), tc.sis("Luz", co.INGRESO, 5, 77, "PP - APARTAMENTO 3", "A")])
assert todo["banco_sin_registro"] == [] and todo["sistema_sin_banco"] == []
openpyxl.load_workbook(io.BytesIO(rep.generar_excel(todo)))

# ---- PDF ----
for resultado in (r, sin_ef, vacio, todo):
    assert rep.generar_pdf(resultado, ext, "Vio")[:4] == b"%PDF"
assert rep.generar_pdf(r)[:4] == b"%PDF"
print("PRUEBAS DE REPORTES DE CONCILIACIÓN OK")
