"""
Lectura del extracto bancario (archivo Excel .xlsx o CSV) para la conciliación bancaria.

Las columnas se detectan por el NOMBRE de su encabezado (sin distinguir mayúsculas ni acentos), así que no importa en qué
orden vengan ni si hay filas de título antes de la tabla:

    Fecha | Hora | Nro Trn./Cheque | Transacción | Nota | (Detalle) | Monto | Saldo

  * fecha, monto (con signo: negativo = débito) o bien dos columnas «Débito» y «Crédito», son obligatorias;
  * hora, nro de transacción, tipo de transacción, nota (a quién o de quién), detalle (una columna sin título entre la nota
    y el monto, o titulada «Detalle»/«Observación», donde se anota de qué fue el movimiento) y saldo son opcionales.
  * Las fechas pueden venir como fecha de Excel o como texto («01/Oct/2026», «01/10/2026», «2026-10-01»).
  * Los montos pueden venir como número o como texto («1,100.00», «1.100,00», «(150.00)», «150.00-», «Bs 150»).

También controla que el SALDO del extracto sea coherente (saldo anterior + monto = saldo), para detectar líneas que faltan,
sobran o están repetidas.

Módulo sin dependencias de Streamlit ni de la base. Lee .xlsx con openpyxl y .csv con el módulo estándar.
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime, time, timedelta

import openpyxl

_MESES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12}

# En este orden se asigna cada columna (así «Nro Trn./Cheque» no se confunde con «Transacción»).
CAMPOS = ["fecha", "hora", "saldo", "monto", "debito", "credito", "transaccion", "nro", "contraparte", "detalle"]
SINONIMOS = {
    "fecha": ("fecha",),
    "hora": ("hora",),
    "saldo": ("saldo",),
    "monto": ("monto", "importe"),
    "debito": ("debito", "debe", "retiro", "retiros", "egreso", "cargo"),
    "credito": ("credito", "haber", "deposito", "depositos", "ingreso", "abono"),
    "transaccion": ("transaccion", "transacciones", "operacion", "movimiento", "tipo"),
    "nro": ("nro", "numero", "trn", "cheque", "referencia", "comprobante", "documento", "secuencia"),
    "contraparte": ("nota", "glosa", "descripcion", "beneficiario", "ordenante", "nombre", "cliente"),
    "detalle": ("detalle", "observacion", "observaciones", "comentario", "comentarios", "concepto"),
}
NOMBRES_CAMPOS = {"fecha": "Fecha", "hora": "Hora", "saldo": "Saldo", "monto": "Monto", "debito": "Débito", "credito": "Crédito",
                  "transaccion": "Transacción", "nro": "N° de transacción", "contraparte": "Nota (a quién / de quién)",
                  "detalle": "Detalle (tu anotación)"}
TOLERANCIA_SALDO = 0.015


class ErrorExtracto(Exception):
    """El archivo no se pudo leer; el mensaje está en español y se puede mostrar tal cual."""


def _norm(texto):
    t = unicodedata.normalize("NFD", str(texto if texto is not None else ""))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn").casefold()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", t).split())


# ============================================================
# Conversión de valores
# ============================================================

def parsear_fecha(valor):
    """date a partir de una fecha de Excel o de un texto; None si no se puede."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        if 20000 < valor < 80000:                                    # número de serie de Excel
            return (datetime(1899, 12, 30) + timedelta(days=float(valor))).date()
        return None
    texto = str(valor).strip() if valor is not None else ""
    if not texto:
        return None
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", texto)                           # 2026-10-01
    if m:
        return _crear_fecha(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"^(\d{1,2})[\s/.\-]+([A-Za-zÁÉÍÓÚáéíóú]{3,10})\.?[\s/.\-]+(\d{2,4})", texto)  # 01/Oct/2026
    if m:
        mes = _MESES.get(_norm(m.group(2))[:3])
        return _crear_fecha(_anio(int(m.group(3))), mes, int(m.group(1))) if mes else None
    m = re.match(r"^(\d{1,2})[\s/.\-]+(\d{1,2})[\s/.\-]+(\d{2,4})", texto)                  # 01/10/2026 (día/mes/año)
    if m:
        return _crear_fecha(_anio(int(m.group(3))), int(m.group(2)), int(m.group(1)))
    return None


def _anio(a):
    return a + 2000 if a < 100 else a


def _crear_fecha(anio, mes, dia):
    try:
        return date(anio, mes, dia)
    except (TypeError, ValueError):
        return None


def parsear_hora(valor):
    if isinstance(valor, datetime):
        return valor.strftime("%H:%M:%S")
    if isinstance(valor, time):
        return valor.strftime("%H:%M:%S")
    return (str(valor).strip() or None) if valor is not None else None


def parsear_monto(valor):
    """float a partir de un número o un texto con formato de dinero; None si no se puede (o está vacío)."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip()
    if not texto:
        return None
    negativo = False
    if texto.startswith("(") and texto.endswith(")"):
        negativo, texto = True, texto[1:-1]
    if texto.endswith("-"):
        negativo, texto = True, texto[:-1]
    texto = re.sub(r"(?i)(?<![a-z])(?:bs|bob|usd|us)(?![a-z])\.?|\$", "", texto).replace(" ", "").replace("\xa0", "")   # quita «Bs», «BOB», «USD», «$»
    if texto.startswith("-"):
        negativo, texto = True, texto[1:]
    elif texto.startswith("+"):
        texto = texto[1:]
    if not re.match(r"^[\d.,]+$", texto) or not re.search(r"\d", texto):
        return None
    if "," in texto and "." in texto:                                # el último separador es el decimal
        decimal = "," if texto.rfind(",") > texto.rfind(".") else "."
        miles = "." if decimal == "," else ","
        texto = texto.replace(miles, "").replace(decimal, ".")
    elif "," in texto:
        partes = texto.split(",")
        texto = texto.replace(",", ".") if len(partes) == 2 and len(partes[1]) != 3 else texto.replace(",", "")
    elif texto.count(".") > 1:
        texto = texto.replace(".", "")
    try:
        numero = float(texto)
    except ValueError:
        return None
    return -numero if negativo else numero


# ============================================================
# Detección de columnas
# ============================================================

def _coincide(encabezado, sinonimos):
    palabras = _norm(encabezado).split()
    return any(p == s or (len(s) >= 5 and p.startswith(s)) for p in palabras for s in sinonimos)


def detectar_fila_encabezado(filas, maximo=40):
    """Índice (0) de la fila que parece el encabezado: tiene «fecha» y alguna columna de monto/débito/crédito. Si ninguna fila tiene
    las dos cosas, se usa la primera que tenga «fecha» (la pantalla deja elegir a mano las columnas que falten)."""
    con_fecha = None
    for i, fila in enumerate(filas[:maximo]):
        textos = [c for c in fila if isinstance(c, str)]
        if any(_coincide(t, SINONIMOS["fecha"]) for t in textos):
            if any(_coincide(t, SINONIMOS["monto"] + SINONIMOS["debito"] + SINONIMOS["credito"]) for t in textos):
                return i
            if con_fecha is None:
                con_fecha = i
    return con_fecha


def detectar_columnas(encabezados):
    """{campo: índice de columna} según los nombres de los encabezados."""
    asignado, usadas = {}, set()
    for campo in CAMPOS:
        for i, enc in enumerate(encabezados):
            if i not in usadas and enc not in (None, "") and _coincide(enc, SINONIMOS[campo]):
                asignado[campo] = i
                usadas.add(i)
                break
    # una columna SIN título entre la nota y el monto es la anotación del usuario (detalle)
    if "detalle" not in asignado:
        izq = asignado.get("contraparte", asignado.get("transaccion"))
        der = asignado.get("monto", asignado.get("debito", asignado.get("credito")))
        if izq is not None and der is not None:
            for i in range(izq + 1, der):
                if i not in usadas and encabezados[i] in (None, ""):
                    asignado["detalle"] = i
                    break
    return asignado


def _celda(fila, mapeo, campo):
    i = mapeo.get(campo)
    return fila[i] if i is not None and i < len(fila) else None


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


# ============================================================
# Lectura del archivo
# ============================================================

def _filas_xlsx(contenido, hoja):
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception:
        raise ErrorExtracto("No se pudo abrir el archivo Excel. Verifica que sea un .xlsx válido (no .xls).") from None
    ws = libro[hoja] if hoja in libro.sheetnames else libro[libro.sheetnames[0]]
    return [list(f) for f in ws.iter_rows(values_only=True)], libro.sheetnames


def _filas_csv(contenido):
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = "\n".join(texto.splitlines()[:15])
    try:
        delimitador = csv.Sniffer().sniff(muestra, delimiters=",;\t|").delimiter
    except csv.Error:
        delimitador = ";" if muestra.count(";") > muestra.count(",") else ","
    return [list(f) for f in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def leer_filas(contenido, nombre, hoja=None):
    """(filas, nombres de hojas) del archivo, sin interpretar."""
    nombre = (nombre or "").lower()
    if nombre.endswith(".xls"):
        raise ErrorExtracto("El formato .xls (Excel antiguo) no se puede leer. Ábrelo en Excel y guárdalo como .xlsx o como .csv.")
    if nombre.endswith((".csv", ".txt")):
        return _filas_csv(contenido), []
    return _filas_xlsx(contenido, hoja)


def leer_extracto(contenido, nombre, mapeo=None, hoja=None):
    """Lee el extracto. mapeo: {campo: índice de columna} para forzar las columnas si la detección falla.
    Devuelve {'lineas', 'columnas', 'encabezados', 'fila_encabezado', 'avisos', 'desde', 'hasta', 'total_creditos', 'total_debitos',
    'saldo_inicial', 'saldo_final', 'problemas_saldo', 'hojas', 'faltan'}. Si faltan columnas obligatorias, 'lineas' queda vacía y
    'faltan' dice cuáles (la pantalla deja elegirlas a mano)."""
    filas, hojas = leer_filas(contenido, nombre, hoja)
    filas = [f for f in filas if any(c not in (None, "") for c in f)]
    if not filas:
        raise ErrorExtracto("El archivo está vacío.")
    i_enc = detectar_fila_encabezado(filas)
    if i_enc is None and not mapeo:
        raise ErrorExtracto("No encontré la fila de encabezados. El extracto debe tener, al menos, una columna «Fecha» y una «Monto» (o «Débito» y «Crédito»).")
    i_enc = i_enc if i_enc is not None else 0
    encabezados = [_texto(c) for c in filas[i_enc]]
    columnas = dict(mapeo) if mapeo else detectar_columnas(filas[i_enc])
    resultado = {"lineas": [], "columnas": columnas, "encabezados": encabezados, "fila_encabezado": i_enc + 1, "avisos": [], "hojas": hojas,
                 "desde": None, "hasta": None, "total_creditos": 0.0, "total_debitos": 0.0, "saldo_inicial": None, "saldo_final": None,
                 "problemas_saldo": [], "faltan": []}
    if "fecha" not in columnas:
        resultado["faltan"].append("fecha")
    if "monto" not in columnas and not ("debito" in columnas or "credito" in columnas):
        resultado["faltan"].append("monto")
    if resultado["faltan"]:
        return resultado

    lineas, ignoradas = [], 0
    for numero, fila in enumerate(filas[i_enc + 1:], start=i_enc + 2):
        fecha = parsear_fecha(_celda(fila, columnas, "fecha"))
        if "monto" in columnas:
            monto = parsear_monto(_celda(fila, columnas, "monto"))
        else:
            cred, deb = parsear_monto(_celda(fila, columnas, "credito")), parsear_monto(_celda(fila, columnas, "debito"))
            monto = None if cred is None and deb is None else (cred or 0.0) - abs(deb or 0.0)
        if fecha is None or monto is None:
            ignoradas += 1                                          # título repetido, fila de totales, etc.
            continue
        lineas.append({"n": len(lineas) + 1, "fila": numero, "fecha": fecha, "hora": parsear_hora(_celda(fila, columnas, "hora")),
                       "nro": _texto(_celda(fila, columnas, "nro")), "transaccion": _texto(_celda(fila, columnas, "transaccion")),
                       "contraparte": _texto(_celda(fila, columnas, "contraparte")), "detalle": _texto(_celda(fila, columnas, "detalle")),
                       "monto": round(monto, 2), "saldo": (round(s, 2) if (s := parsear_monto(_celda(fila, columnas, "saldo"))) is not None else None)})
    if not lineas:
        raise ErrorExtracto("No encontré movimientos con fecha y monto debajo de los encabezados.")

    # Si el banco exporta montos sin signo, el tipo de transacción dice si es débito o crédito
    if "monto" in columnas and all(l["monto"] >= 0 for l in lineas) and "transaccion" in columnas:
        deb = [l for l in lineas if "debito" in _norm(l["transaccion"])]
        if deb and any("credito" in _norm(l["transaccion"]) for l in lineas):
            for l in deb:
                l["monto"] = -abs(l["monto"])
            resultado["avisos"].append("Los montos venían sin signo: se tomaron como débitos los movimientos cuya transacción dice «DÉBITO».")
    if ignoradas:
        resultado["avisos"].append(f"Se ignoraron {ignoradas} fila(s) sin fecha o sin monto (títulos, totales o filas vacías).")
    resultado.update(_resumen(lineas))
    resultado["lineas"] = lineas
    return resultado


def _resumen(lineas):
    fechas = [l["fecha"] for l in lineas]
    # orden cronológico real del archivo: si las fechas bajan, el banco lo exportó del más nuevo al más antiguo
    descendente = len(lineas) > 1 and fechas[0] > fechas[-1]
    orden = list(reversed(lineas)) if descendente else list(lineas)
    problemas, previo = [], None
    for l in orden:
        if previo is not None and l["saldo"] is not None and previo["saldo"] is not None:
            esperado = round(previo["saldo"] + l["monto"], 2)
            if abs(esperado - l["saldo"]) > TOLERANCIA_SALDO:
                problemas.append({"fila": l["fila"], "fecha": l["fecha"], "monto": l["monto"], "saldo_esperado": esperado,
                                  "saldo_extracto": l["saldo"], "diferencia": round(l["saldo"] - esperado, 2)})
        previo = l if l["saldo"] is not None else previo
    primero = next((l for l in orden if l["saldo"] is not None), None)
    ultimo = next((l for l in reversed(orden) if l["saldo"] is not None), None)
    return {"desde": min(fechas), "hasta": max(fechas),
            "total_creditos": round(sum(l["monto"] for l in lineas if l["monto"] > 0), 2),
            "total_debitos": round(-sum(l["monto"] for l in lineas if l["monto"] < 0), 2),
            "saldo_inicial": round(primero["saldo"] - primero["monto"], 2) if primero else None,
            "saldo_final": ultimo["saldo"] if ultimo else None, "problemas_saldo": problemas}
