"""Pruebas de cuadre.py  ->  python test_cuadre.py   (necesita recibo.py con su logo y fuentes al lado)"""
import io
import cuadre as c
import estados

ventas = [
    {"id": 1, "fecha_venta": "2026-10-02", "concepto": "Alquiler de área común", "monto": 300, "forma_cobro": "Efectivo",
     "comprador": "Ana", "recibo_emitido": "R-1", "estado_entrega": "No aplica", "encargado": "Vio"},
    {"id": 2, "fecha_venta": "2026-10-05", "concepto": "Emisión de tag/control de acceso", "monto": 150.5, "forma_cobro": "QR",
     "comprador": "Luis", "recibo_emitido": "", "estado_entrega": "Pendiente de entrega", "encargado": "Vio"},
]
pagos = [
    {"id": 1, "fecha_pago": "2026-10-02", "categoria": "Sueldos y honorarios", "descripcion": "=1+1 sueldo", "monto": 120.25,
     "metodo_pago": "efectivo", "beneficiario": "Juan", "estado_devolucion": "Se realizó devolución", "encargado": "Vio"},   # texto anterior, sin «devolución a»
    {"id": 2, "fecha_pago": "2026-10-09", "categoria": "Servicios (luz, agua, internet)", "descripcion": "Luz", "monto": 80,
     "metodo_pago": "Transferencia Bancaria", "beneficiario": "ENDE", "estado_devolucion": "Pendiente de devolución", "encargado": "Vio"},
]
compras = [
    {"id": 1, "fecha_compra": "2026-10-02", "categoria": "Artículos de limpieza", "descripcion": "Escobas", "monto_total": 45.1,
     "metodo_pago": "Efectivo", "proveedor": "Tienda", "numero_comprobante": "F-9", "estado_devolucion": None, "encargado": "Vio"},
    {"id": 2, "fecha_compra": "2026-10-31", "categoria": "Seguridad", "descripcion": "Cámara", "monto_total": 200,
     "metodo_pago": "", "proveedor": "", "numero_comprobante": None, "estado_devolucion": estados.REALIZADA_DEVOLUCION, "devolucion_a": "Marta", "encargado": "Vio"},
    {"id": 3, "fecha_compra": None, "categoria": "X", "monto_total": 999},     # sin fecha: se ignora
]

# ---- métodos ----
g = c.grupo_metodo
assert g("Efectivo") == g(" EFECTIVO ") == g("efectivo") == "Efectivo"
assert g("QR") == g("qr") == g("Pago QR") == "QR" and g("Transferencia Bancaria") == g("transferencia") == "Transferencia"
assert g("Tarjeta") == "Tarjeta" and g("") == g(None) == "Sin método" and g("Cheque") == "Otro"
assert g("Parqueo") == "Otro"                                    # 'qr' solo cuenta como palabra suelta

movs = c.construir_movimientos(pagos, compras, ventas)
assert len(movs) == 6                                            # el de sin fecha no entra
assert [m["fecha"].isoformat() for m in movs] == sorted(m["fecha"].isoformat() for m in movs)
q = c.construir_cuadre(movs, "2026-10-01", "2026-10-31", saldo_inicial=1000, efectivo_contado=1100)

assert q["ingresos"] == 450.5 and q["egresos"] == 445.35 and q["resultado"] == 5.15      # 120.25+80+45.1+200
assert [o["total"] for o in q["por_origen"]] == [450.5, 200.25, 245.1] and [o["n"] for o in q["por_origen"]] == [2, 2, 2]
met = {m["metodo"]: m for m in q["por_metodo"]}
assert met["Efectivo"]["ingresos"] == 300 and met["Efectivo"]["egresos"] == 165.35      # 120.25 + 45.10
assert met["QR"]["ingresos"] == 150.5 and met["Transferencia"]["egresos"] == 80 and met["Sin método"]["egresos"] == 200
assert round(sum(m["ingresos"] for m in q["por_metodo"]), 2) == q["ingresos"]            # los métodos suman el total
assert round(sum(m["egresos"] for m in q["por_metodo"]), 2) == q["egresos"]
ef = q["efectivo"]
assert ef["esperado"] == 1134.65 and ef["diferencia"] == -34.65 and ef["contado"] == 1100   # 1000 + 300 - 165.35
assert c.construir_cuadre(movs, saldo_inicial=0)["efectivo"]["diferencia"] is None          # sin contar el efectivo: sin diferencia
assert q["libro"][-1]["acumulado"] == q["resultado"] == q["por_dia"][-1]["acumulado"]
assert round(sum(d["neto"] for d in q["por_dia"]), 2) == q["resultado"] and len(q["por_dia"]) == 4
assert q["pendientes"] == {"devolucion_pendiente": 80, "devolucion_realizada": 320.25, "entrega_pendiente": 150.5, "entrega_realizada": 0}
assert q["categorias_egreso"][0] == {"categoria": "Seguridad", "n": 1, "total": 200}       # ordenado de mayor a menor
assert sum(x["total"] for x in q["categorias_egreso"]) == q["egresos"]
motivos = " | ".join(r["Por revisar"] for r in q["revisar"])
assert "Devolución realizada sin indicar a quién" in motivos and "Sin método" in motivos and "Compra sin N° de comprobante" in motivos and "Venta sin N° de recibo" in motivos
assert [f["Comprobante N°"] for f in c.filas_por_origen(q, "Compra")] == ["F-9", ""] and "Comprobante N°" not in c.filas_por_origen(q, "Pago")[0]
vacio = c.construir_cuadre(c.construir_movimientos([], [], []), "2026-10-01", "2026-10-31")
assert vacio["n_mov"] == 0 and vacio["resultado"] == 0 and vacio["libro"] == [] and vacio["efectivo"]["esperado"] == 0

# ---- a qué usuario se devolvió ----
por_id = {(m["origen"], m["id"]): m for m in movs}
assert por_id[("Compra", 2)]["devolucion_a"] == "Marta"                                   # realizada, con nombre
assert por_id[("Pago", 1)]["devolucion_a"] == "(sin indicar)" and por_id[("Pago", 1)]["estado"] == estados.REALIZADA_DEVOLUCION   # texto anterior = realizada
assert por_id[("Pago", 2)]["devolucion_a"] == "" and por_id[("Compra", 1)]["devolucion_a"] == ""                                 # pendiente / no aplica: vacío
assert all(m["devolucion_a"] == "" for m in movs if m["origen"] == "Venta")                                                      # las ventas no tienen devolución
assert [f["Devolución a"] for f in c.filas_por_origen(q, "Compra")] == ["", "Marta"] and "Devolución a" not in c.filas_por_origen(q, "Venta")[0]
assert q["libro"][0]["devolucion_a"] in ("", "(sin indicar)", "Marta")

# ---- Excel ----
import openpyxl
datos = c.generar_excel_cuadre(q, "Vio")
wb = openpyxl.load_workbook(io.BytesIO(datos))
assert wb.sheetnames[:4] == ["Resumen", "Libro diario", "Por día", "Por método"] and "Por revisar" in wb.sheetnames
libro = wb["Libro diario"]
descripciones = [libro.cell(r, 5).value for r in range(5, 11)]
assert "=1+1 sueldo" in descripciones                                                      # texto, no fórmula
assert all(cell.data_type != "f" for row in libro.iter_rows() for cell in row)
assert libro.cell(11, 11).value == 450.5 and libro.cell(11, 12).value == 445.35            # fila TOTAL
# vacío no debe fallar
openpyxl.load_workbook(io.BytesIO(c.generar_excel_cuadre(vacio)))

# ---- PDF ----
pdf = c.generar_pdf_cuadre(q, "Vio")
assert pdf[:4] == b"%PDF" and len(pdf) > 3000
assert c.generar_pdf_cuadre(vacio)[:4] == b"%PDF"
print("PRUEBAS DE CUADRE OK")
