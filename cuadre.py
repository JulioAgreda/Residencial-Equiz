"""
Cuadre de tesorería: libro de movimientos del periodo (Pagos, Compras y Ventas) con los totales que necesita
la tesorera para cuadrar la caja a fin de mes.

  * Ingresos = Ventas.   Egresos = Pagos + Compras.   Resultado = Ingresos - Egresos.
  * Los montos NO se descuentan según el «Estado» (devolución / entrega): ese dato se muestra aparte para que la
    tesorera decida cómo tratarlo.
  * Cuadre de efectivo: saldo inicial + ingresos en efectivo - egresos en efectivo = efectivo esperado en caja;
    se compara con el efectivo realmente contado.

Módulo sin dependencias de Streamlit ni de la base: recibe los registros de la base (diccionarios) y devuelve
datos listos para mostrar o exportar a Excel y PDF.
"""
import io
import os
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (CondPageBreak, HRFlowable, Image as RLImage, KeepTogether, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

import estados
import recibo
import reportes as rp

INGRESO, EGRESO = "Ingreso", "Egreso"
EFECTIVO, QR, TRANSFERENCIA, TARJETA, OTRO, SIN_METODO = "Efectivo", "QR", "Transferencia", "Tarjeta", "Otro", "Sin método"
ORDEN_METODOS = [EFECTIVO, QR, TRANSFERENCIA, TARJETA, OTRO, SIN_METODO]
_ORDEN_ORIGEN = {"Venta": 0, "Pago": 1, "Compra": 2}


# ============================================================
# Utilidades
# ============================================================

def _d(valor):
    """Decimal exacto de 2 decimales (evita errores de coma flotante al sumar montos)."""
    try:
        return Decimal(str(round(float(valor or 0), 2)))
    except (TypeError, ValueError):
        return Decimal(0)


def _f(decimal):
    return float(round(decimal, 2))


def _fecha(valor):
    if isinstance(valor, date):
        return valor
    try:
        return date.fromisoformat(str(valor)[:10]) if valor else None
    except ValueError:
        return None


def grupo_metodo(valor):
    """Agrupa el método de pago/cobro (texto libre en algunos registros) en: Efectivo, QR, Transferencia,
    Tarjeta, Otro o Sin método."""
    t = unicodedata.normalize("NFD", str(valor or ""))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn").casefold().strip()
    if not t:
        return SIN_METODO
    palabras = set(re.split(r"[^a-z0-9]+", t))
    if "efectivo" in t:
        return EFECTIVO
    if "qr" in palabras:
        return QR
    if "transfer" in t or "deposito" in t:
        return TRANSFERENCIA
    if "tarjeta" in t or "debito" in t or "credito" in t:
        return TARJETA
    return OTRO


# ============================================================
# Datos: registros de la base -> movimientos unificados
# ============================================================

def construir_movimientos(pagos, compras, ventas):
    """Lista única de movimientos, ordenada por fecha. Cada uno trae: fecha, tipo (Ingreso/Egreso), origen
    (Venta/Pago/Compra), categoría, detalle, persona, comprobante, método (agrupado), estado y monto."""
    movs = []

    def agregar(r, origen, tipo, campo_fecha, campo_monto, categoria, persona, comprobante, campo_metodo, campo_estado, opciones,
                con_devolucion=False):
        fecha = _fecha(r.get(campo_fecha))
        if not fecha:
            return
        estado = estados.normalizar(r.get(campo_estado), opciones)
        movs.append({
            "id": r.get("id") or 0, "fecha": fecha, "tipo": tipo, "origen": origen,
            "categoria": rp._txt(r.get(categoria)), "detalle": rp._txt(r.get("descripcion")),
            "persona": rp._txt(r.get(persona)), "comprobante": rp._txt(r.get(comprobante)),
            "metodo": grupo_metodo(r.get(campo_metodo)), "metodo_original": rp._txt(r.get(campo_metodo)),
            "estado": estado, "monto": _f(_d(r.get(campo_monto))), "encargado": rp._txt(r.get("encargado")),
            # a quién se le devolvió el dinero: usuario o «Tercero: nombre» (solo Pagos y Compras con la devolución ya realizada)
            "devolucion_a": estados.texto_devolucion_a(estado, r.get("devolucion_a")) if con_devolucion else "",
        })

    for r in ventas or []:
        agregar(r, "Venta", INGRESO, "fecha_venta", "monto", "concepto", "comprador", "recibo_emitido",
                "forma_cobro", "estado_entrega", estados.ESTADOS_ENTREGA)
    for r in pagos or []:
        agregar(r, "Pago", EGRESO, "fecha_pago", "monto", "categoria", "beneficiario", None,
                "metodo_pago", "estado_devolucion", estados.ESTADOS_DEVOLUCION, con_devolucion=True)
    for r in compras or []:
        agregar(r, "Compra", EGRESO, "fecha_compra", "monto_total", "categoria", "proveedor", "numero_comprobante",
                "metodo_pago", "estado_devolucion", estados.ESTADOS_DEVOLUCION, con_devolucion=True)
    movs.sort(key=lambda m: (m["fecha"], _ORDEN_ORIGEN[m["origen"]], m["id"]))
    return movs


def _agrupar_categorias(movs, tipo):
    acum = {}
    for m in movs:
        if m["tipo"] == tipo:
            nombre = m["categoria"] or "(Sin categoría)"
            n, total = acum.get(nombre, (0, Decimal(0)))
            acum[nombre] = (n + 1, total + _d(m["monto"]))
    filas = [{"categoria": k, "n": n, "total": _f(t)} for k, (n, t) in acum.items()]
    filas.sort(key=lambda f: (-f["total"], f["categoria"].casefold()))
    return filas


def construir_cuadre(movs, desde=None, hasta=None, saldo_inicial=0, efectivo_contado=None):
    """Todo lo que muestra y exporta el reporte, calculado a partir de los movimientos."""
    ingresos = sum((_d(m["monto"]) for m in movs if m["tipo"] == INGRESO), Decimal(0))
    egresos = sum((_d(m["monto"]) for m in movs if m["tipo"] == EGRESO), Decimal(0))

    por_origen = []
    for origen, titulo, tipo in (("Venta", "Ventas", INGRESO), ("Pago", "Pagos", EGRESO), ("Compra", "Compras", EGRESO)):
        sel = [m for m in movs if m["origen"] == origen]
        por_origen.append({"origen": titulo, "tipo": tipo, "n": len(sel),
                           "total": _f(sum((_d(m["monto"]) for m in sel), Decimal(0)))})

    por_metodo = []
    for metodo in ORDEN_METODOS:
        sel = [m for m in movs if m["metodo"] == metodo]
        if not sel:
            continue
        ing = sum((_d(m["monto"]) for m in sel if m["tipo"] == INGRESO), Decimal(0))
        egr = sum((_d(m["monto"]) for m in sel if m["tipo"] == EGRESO), Decimal(0))
        por_metodo.append({"metodo": metodo, "n": len(sel), "ingresos": _f(ing), "egresos": _f(egr), "neto": _f(ing - egr)})

    dias, acumulado = {}, Decimal(0)
    for m in movs:
        d = dias.setdefault(m["fecha"], {"ingresos": Decimal(0), "egresos": Decimal(0), "n": 0})
        d["ingresos" if m["tipo"] == INGRESO else "egresos"] += _d(m["monto"])
        d["n"] += 1
    por_dia = []
    for fecha in sorted(dias):
        d = dias[fecha]
        acumulado += d["ingresos"] - d["egresos"]
        por_dia.append({"fecha": fecha, "fecha_txt": rp._fmt_fecha(fecha), "n": d["n"], "ingresos": _f(d["ingresos"]),
                        "egresos": _f(d["egresos"]), "neto": _f(d["ingresos"] - d["egresos"]), "acumulado": _f(acumulado)})

    libro, acumulado = [], Decimal(0)
    for m in movs:
        monto = _d(m["monto"])
        acumulado += monto if m["tipo"] == INGRESO else -monto
        libro.append({**m, "fecha_txt": rp._fmt_fecha(m["fecha"]),
                      "ingreso": m["monto"] if m["tipo"] == INGRESO else None,
                      "egreso": m["monto"] if m["tipo"] == EGRESO else None, "acumulado": _f(acumulado)})

    ing_ef = sum((_d(m["monto"]) for m in movs if m["tipo"] == INGRESO and m["metodo"] == EFECTIVO), Decimal(0))
    egr_ef = sum((_d(m["monto"]) for m in movs if m["tipo"] == EGRESO and m["metodo"] == EFECTIVO), Decimal(0))
    inicial = _d(saldo_inicial)
    esperado = inicial + ing_ef - egr_ef
    efectivo = {"saldo_inicial": _f(inicial), "ingresos": _f(ing_ef), "egresos": _f(egr_ef), "esperado": _f(esperado),
                "contado": None if efectivo_contado is None else _f(_d(efectivo_contado)),
                "diferencia": None if efectivo_contado is None else _f(_d(efectivo_contado) - esperado)}

    def suma_estado(tipo, estado):
        return _f(sum((_d(m["monto"]) for m in movs if m["tipo"] == tipo and m["estado"] == estado), Decimal(0)))

    pendientes = {
        "devolucion_pendiente": suma_estado(EGRESO, estados.PENDIENTE_DEVOLUCION),
        "devolucion_realizada": suma_estado(EGRESO, estados.REALIZADA_DEVOLUCION),
        "entrega_pendiente": suma_estado(INGRESO, estados.PENDIENTE_ENTREGA),
        "entrega_realizada": suma_estado(INGRESO, estados.ESTADOS_ENTREGA[2]),
    }

    revisar = []
    for m in movs:
        problemas = []
        if m["metodo"] == SIN_METODO:
            problemas.append("Sin método de pago/cobro")
        if m["monto"] <= 0:
            problemas.append("Monto en cero")
        if m["origen"] == "Compra" and not m["comprobante"]:
            problemas.append("Compra sin N° de comprobante")
        if m["origen"] == "Venta" and not m["comprobante"]:
            problemas.append("Venta sin N° de recibo")
        if m["devolucion_a"] == estados.SIN_INDICAR:
            problemas.append("Devolución realizada sin indicar a quién")
        if problemas:
            revisar.append({"Fecha": rp._fmt_fecha(m["fecha"]), "Registro": m["origen"],
                            "Categoría": m["categoria"] or "—", "Detalle": m["detalle"] or "—",
                            "Monto": m["monto"], "Por revisar": "; ".join(problemas)})

    return {"desde": desde, "hasta": hasta, "n_mov": len(movs), "ingresos": _f(ingresos), "egresos": _f(egresos),
            "resultado": _f(ingresos - egresos), "por_origen": por_origen, "por_metodo": por_metodo, "por_dia": por_dia,
            "categorias_ingreso": _agrupar_categorias(movs, INGRESO), "categorias_egreso": _agrupar_categorias(movs, EGRESO),
            "libro": libro, "efectivo": efectivo, "pendientes": pendientes, "revisar": revisar}


def filas_por_origen(cuadre, origen):
    """Detalle (como el de «Actividad por usuario») de 'Venta', 'Pago' o 'Compra' para mostrar o exportar."""
    out = []
    for m in cuadre["libro"]:
        if m["origen"] != origen:
            continue
        fila = {"Fecha": m["fecha_txt"], "Categoría" if origen != "Venta" else "Concepto": m["categoria"],
                "Descripción": m["detalle"],
                {"Venta": "Comprador", "Pago": "Beneficiario", "Compra": "Proveedor"}[origen]: m["persona"]}
        if origen != "Pago":
            fila["Recibo N°" if origen == "Venta" else "Comprobante N°"] = m["comprobante"]
        fila["Método"] = m["metodo_original"] or "—"
        fila["Estado"] = rp._estado_txt(m["estado"])
        if origen != "Venta":
            fila["Devolución a"] = m["devolucion_a"]
        fila["Monto"] = m["monto"]
        out.append(fila)
    return out


# ============================================================
# EXCEL
# ============================================================

_FMT = rp._FMT_EXCEL


def _hoja(wb, nombre, titulo, periodo, columnas, filas, totales=None):
    """Hoja con una tabla. columnas: [(clave, encabezado, 'txt'|'bs'|'num')]. totales: {clave: valor}."""
    ws = wb.create_sheet(nombre)
    rp._poner(ws, 1, 1, titulo, Font(name="Calibri", bold=True, size=14))
    rp._poner(ws, 2, 1, periodo, Font(name="Calibri", size=10))
    f_header = Font(name="Calibri", bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor=rp._COLOR_HEADER)
    fila0 = 4
    anchos = {}
    for col, (_c, enc, _t) in enumerate(columnas, start=1):
        c = rp._poner(ws, fila0, col, enc, f_header)
        c.fill = relleno
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        anchos[col] = len(enc)
    fila = fila0 + 1
    for f in filas:
        for col, (clave, _e, tipo) in enumerate(columnas, start=1):
            valor = f.get(clave)
            c = rp._poner(ws, fila, col, valor, Font(name="Calibri", size=10), _FMT if tipo == "bs" else None)
            if tipo == "txt" and clave == "detalle":
                c.alignment = Alignment(wrap_text=True, vertical="top")
            anchos[col] = max(anchos[col], 14 if tipo != "txt" else min(len(str(valor or "")), 45))
        fila += 1
    if totales:
        f_total = Font(name="Calibri", bold=True, size=11)
        rp._poner(ws, fila, 1, "TOTAL", f_total)
        for col, (clave, _e, _t) in enumerate(columnas, start=1):
            if clave in totales:
                rp._poner(ws, fila, col, totales[clave], f_total, _FMT if columnas[col - 1][2] == "bs" else None)
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = max(12, min(48, ancho + 2))
    ws.freeze_panes = ws.cell(fila0 + 1, 1)
    return ws


def generar_excel_cuadre(cuadre, emitido_por=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    periodo = f"Periodo: {rp._fmt_fecha(cuadre['desde'])} al {rp._fmt_fecha(cuadre['hasta'])}"
    f_titulo = Font(name="Calibri", bold=True, size=14)
    f_normal = Font(name="Calibri", size=10)
    f_bold = Font(name="Calibri", bold=True, size=11)
    f_sec = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    relleno = PatternFill("solid", fgColor=rp._COLOR_HEADER)

    rp._poner(ws, 1, 1, "Cuadre de tesorería", f_titulo)
    rp._poner(ws, 2, 1, periodo, f_normal)
    if emitido_por:
        rp._poner(ws, 3, 1, f"Emitido por: {emitido_por}", f_normal)

    fila = [5]

    def seccion(texto):
        c = rp._poner(ws, fila[0], 1, texto, f_sec)
        for col in (1, 2, 3, 4):
            ws.cell(fila[0], col).fill = relleno
        fila[0] += 1

    def linea(etiqueta, valor, negrita=False, formato=_FMT):
        rp._poner(ws, fila[0], 1, etiqueta, f_bold if negrita else f_normal)
        if valor is not None:
            rp._poner(ws, fila[0], 2, valor, f_bold if negrita else f_normal, formato)
        else:
            rp._poner(ws, fila[0], 2, "—", f_normal)
        fila[0] += 1

    seccion("Totales del periodo")
    linea("Ingresos (ventas)", cuadre["ingresos"], True)
    linea("Egresos (pagos + compras)", cuadre["egresos"], True)
    linea("Resultado (ingresos - egresos)", cuadre["resultado"], True)
    linea("Cantidad de movimientos", cuadre["n_mov"], formato="0")
    fila[0] += 1
    seccion("Por tipo de registro")
    for o in cuadre["por_origen"]:
        linea(f"{o['origen']} ({o['n']})", o["total"])
    fila[0] += 1
    ef = cuadre["efectivo"]
    seccion("Cuadre de efectivo en caja")
    linea("Saldo inicial en efectivo", ef["saldo_inicial"])
    linea("+ Ingresos en efectivo", ef["ingresos"])
    linea("- Egresos en efectivo", ef["egresos"])
    linea("= Efectivo esperado en caja", ef["esperado"], True)
    linea("Efectivo contado", ef["contado"])
    linea("Diferencia (contado - esperado)", ef["diferencia"], True)
    fila[0] += 1
    seccion("Por método de pago / cobro")
    rp._poner(ws, fila[0], 1, "Método", f_bold)
    for col, t in ((2, "Ingresos"), (3, "Egresos"), (4, "Neto")):
        rp._poner(ws, fila[0], col, t, f_bold)
    fila[0] += 1
    for m in cuadre["por_metodo"]:
        rp._poner(ws, fila[0], 1, m["metodo"], f_normal)
        for col, k in ((2, "ingresos"), (3, "egresos"), (4, "neto")):
            rp._poner(ws, fila[0], col, m[k], f_normal, _FMT)
        fila[0] += 1
    fila[0] += 1
    p = cuadre["pendientes"]
    seccion("Según el Estado de cada registro (no se descuenta de los totales)")
    linea("Egresos con devolución pendiente", p["devolucion_pendiente"])
    linea("Egresos con devolución realizada", p["devolucion_realizada"])
    linea("Ingresos con entrega pendiente", p["entrega_pendiente"])
    linea("Ingresos con entrega realizada", p["entrega_realizada"])
    ws.column_dimensions["A"].width = 46
    for col in "BCD":
        ws.column_dimensions[col].width = 20

    _hoja(wb, "Libro diario", "Libro diario de movimientos", periodo,
          [("fecha_txt", "Fecha", "txt"), ("tipo", "Tipo", "txt"), ("origen", "Registro", "txt"),
           ("categoria", "Categoría / Concepto", "txt"), ("detalle", "Descripción", "txt"),
           ("persona", "Persona / Empresa", "txt"), ("comprobante", "N° comprobante / recibo", "txt"),
           ("metodo_original", "Método", "txt"), ("estado", "Estado", "txt"), ("devolucion_a", "Devolución realizada a", "txt"),
           ("ingreso", "Ingreso", "bs"), ("egreso", "Egreso", "bs"), ("acumulado", "Saldo acumulado", "bs")],
          [{**m, "estado": rp._estado_txt(m["estado"])} for m in cuadre["libro"]],
          {"ingreso": cuadre["ingresos"], "egreso": cuadre["egresos"], "acumulado": cuadre["resultado"]})
    _hoja(wb, "Por día", "Totales por día (para cuadres diarios)", periodo,
          [("fecha_txt", "Fecha", "txt"), ("n", "Movimientos", "num"), ("ingresos", "Ingresos", "bs"),
           ("egresos", "Egresos", "bs"), ("neto", "Neto del día", "bs"), ("acumulado", "Acumulado", "bs")],
          cuadre["por_dia"], {"ingresos": cuadre["ingresos"], "egresos": cuadre["egresos"], "neto": cuadre["resultado"]})
    _hoja(wb, "Por método", "Totales por método de pago / cobro", periodo,
          [("metodo", "Método", "txt"), ("n", "Movimientos", "num"), ("ingresos", "Ingresos", "bs"),
           ("egresos", "Egresos", "bs"), ("neto", "Neto", "bs")],
          cuadre["por_metodo"], {"ingresos": cuadre["ingresos"], "egresos": cuadre["egresos"], "neto": cuadre["resultado"]})
    _hoja(wb, "Egresos por categoría", "Egresos por categoría (pagos + compras)", periodo,
          [("categoria", "Categoría", "txt"), ("n", "Movimientos", "num"), ("total", "Total", "bs")],
          cuadre["categorias_egreso"], {"total": cuadre["egresos"]})
    _hoja(wb, "Ingresos por concepto", "Ingresos por concepto (ventas)", periodo,
          [("categoria", "Concepto", "txt"), ("n", "Movimientos", "num"), ("total", "Total", "bs")],
          cuadre["categorias_ingreso"], {"total": cuadre["ingresos"]})
    for origen, nombre, titulo, total in (("Venta", "Ventas", "Ventas realizadas", cuadre["por_origen"][0]["total"]),
                                          ("Pago", "Pagos", "Pagos realizados", cuadre["por_origen"][1]["total"]),
                                          ("Compra", "Compras", "Compras realizadas", cuadre["por_origen"][2]["total"])):
        filas = filas_por_origen(cuadre, origen)
        columnas = [(k, k, "bs" if k == "Monto" else "txt") for k in (filas[0].keys() if filas else ["Fecha", "Monto"])]
        _hoja(wb, nombre, titulo, periodo, columnas, filas, {"Monto": total})
    if cuadre["revisar"]:
        _hoja(wb, "Por revisar", "Registros que conviene revisar antes de cerrar el mes", periodo,
              [(k, k, "bs" if k == "Monto" else "txt") for k in cuadre["revisar"][0].keys()], cuadre["revisar"])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ============================================================
# PDF (horizontal, mismo estilo que los demás reportes)
# ============================================================

def estilos_pdf():
    """Estilos de párrafo comunes a los reportes en PDF horizontales (cuadre, rendición)."""
    from types import SimpleNamespace
    base = getSampleStyleSheet()["Normal"]
    cel = ParagraphStyle("cel", parent=base, fontSize=7, leading=8.5)
    h = ParagraphStyle("h", parent=cel, fontName="Helvetica-Bold", textColor=colors.white)
    return SimpleNamespace(
        titulo=ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=18),
        sub=ParagraphStyle("s", parent=base, fontSize=9, leading=12, textColor=colors.Color(.33, .33, .33)),
        chico=ParagraphStyle("c", parent=base, fontSize=8, leading=10), cel=cel,
        der=ParagraphStyle("d", parent=cel, alignment=2), h=h, hd=ParagraphStyle("hd", parent=h, alignment=2),
        sec=ParagraphStyle("sec", parent=base, fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=8, spaceAfter=3),
        nota=ParagraphStyle("n", parent=base, fontName="Helvetica-Oblique", fontSize=7.5, textColor=colors.Color(.33, .33, .33)))


def documento_pdf(buffer, titulo):
    """Documento horizontal tamaño carta con márgenes de 12 mm. Devuelve (documento, ancho útil)."""
    pagina = landscape(letter)
    doc = SimpleDocTemplate(buffer, pagesize=pagina, topMargin=12 * mm, bottomMargin=12 * mm,
                            leftMargin=12 * mm, rightMargin=12 * mm, title=titulo)
    return doc, pagina[0] - 24 * mm


def tabla_pdf(columnas, filas, pesos, alinear_der=(), fila_total=None, ancho_tabla=None, *, ancho, est):
    """Tabla con encabezado azul. columnas: encabezados; filas: listas de textos; alinear_der: índices de columnas
    numéricas; pesos: ancho relativo de cada columna; ancho_tabla: ancho de esta tabla (por defecto, todo el ancho útil)."""
    pdf = rp._pdf_texto
    datos = [[Paragraph(pdf(c), est.hd if i in alinear_der else est.h) for i, c in enumerate(columnas)]]
    for f in filas:
        datos.append([Paragraph(pdf(v), est.der if i in alinear_der else est.cel) for i, v in enumerate(f)])
    if fila_total:
        datos.append([Paragraph(f"<b>{pdf(v)}</b>", est.der if i in alinear_der else est.cel) for i, v in enumerate(fila_total)])
    t = Table(datos, repeatRows=1, hAlign="LEFT", colWidths=[(ancho_tabla or ancho) * p / sum(pesos) for p in pesos])
    estilo = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + rp._COLOR_HEADER)),
              ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
              ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
    if fila_total:
        estilo += [("LINEABOVE", (0, -1), (-1, -1), 1, colors.black), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f0f0f0"))]
    t.setStyle(TableStyle(estilo))
    return t


def cabecera_pdf(subtitulo, desde, hasta, emitido_por, ancho, est, texto_periodo=None):
    """Logo, nombre del residencial, título del reporte, periodo, fecha de emisión y quién lo emitió.
    texto_periodo: reemplaza «Periodo: desde al hasta» (p. ej. «Al 13/10/2026» en un reporte a una fecha)."""
    pdf = rp._pdf_texto
    logo = ""
    try:
        if os.path.exists(recibo.LOGO_PATH):
            logo = RLImage(io.BytesIO(recibo._logo_liviano_bytes()), width=16 * mm, height=16 * mm)
    except Exception:
        logo = ""
    cab = Table([[logo, [Paragraph(recibo.NOMBRE_RESIDENCIAL, est.titulo), Paragraph(subtitulo, est.sub)]]],
                colWidths=[20 * mm, ancho - 20 * mm])
    cab.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    emitido = datetime.now(recibo._ZONA_BOLIVIA).strftime("%d/%m/%Y %H:%M")
    periodo = (f"<b>Periodo:</b> {rp._fmt_fecha(desde)} al {rp._fmt_fecha(hasta)}" if texto_periodo is None
               else f"<b>{pdf(texto_periodo)}</b>")
    return [cab, Spacer(1, 3), HRFlowable(width="100%", thickness=0.6, color=colors.black), Spacer(1, 4),
            Paragraph(periodo + f" &nbsp;&nbsp;|&nbsp;&nbsp; <b>Emitido:</b> {emitido}"
                      + (f" &nbsp;&nbsp;|&nbsp;&nbsp; <b>Por:</b> {pdf(emitido_por)}" if emitido_por else ""), est.chico),
            Spacer(1, 6)]


def firmas_pdf(ancho, izquierda="Elaborado por (Tesorería)", derecha="Revisado por (Administración)"):
    """Espacio y líneas de firma al final del reporte."""
    firmas = Table([["_______________________________", "", "_______________________________"], [izquierda, "", derecha]],
                   colWidths=[ancho * 0.35, ancho * 0.30, ancho * 0.35])
    firmas.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    return [CondPageBreak(30 * mm), KeepTogether([Spacer(1, 22 * mm), firmas])]


def generar_pdf_cuadre(cuadre, emitido_por=None):
    bs = rp._fmt_bs
    pdf = rp._pdf_texto
    buffer = io.BytesIO()
    doc, ancho = documento_pdf(buffer, "Cuadre de tesorería")
    est = estilos_pdf()
    st_sec, st_nota = est.sec, est.nota

    def tabla(*args, **kwargs):
        return tabla_pdf(*args, ancho=ancho, est=est, **kwargs)

    story = cabecera_pdf("Cuadre de tesorería", cuadre["desde"], cuadre["hasta"], emitido_por, ancho, est)

    # ---- Totales y cuadre de efectivo, lado a lado ----
    ef = cuadre["efectivo"]
    totales = tabla(["Totales del periodo", "Bs"], [
        ["Ingresos (ventas)", bs(cuadre["ingresos"])], ["Egresos (pagos + compras)", bs(cuadre["egresos"])]],
        [3, 2], (1,), ["Resultado (ingresos - egresos)", bs(cuadre["resultado"])], ancho_tabla=ancho * 0.44)
    cuadre_ef = tabla(["Cuadre de efectivo en caja", "Bs"], [
        ["Saldo inicial en efectivo", bs(ef["saldo_inicial"])], ["+ Ingresos en efectivo", bs(ef["ingresos"])],
        ["- Egresos en efectivo", bs(ef["egresos"])], ["= Efectivo esperado en caja", bs(ef["esperado"])],
        ["Efectivo contado", bs(ef["contado"]) if ef["contado"] is not None else "—"]],
        [3, 2], (1,), ["Diferencia (contado - esperado)", bs(ef["diferencia"]) if ef["diferencia"] is not None else "—"],
        ancho_tabla=ancho * 0.50)
    lado = Table([[totales, cuadre_ef]], colWidths=[ancho * 0.46, ancho * 0.54], hAlign="LEFT")
    lado.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [lado]

    story.append(Paragraph("Por método de pago / cobro", st_sec))
    story.append(tabla(["Método", "Movimientos", "Ingresos", "Egresos", "Neto"],
                       [[m["metodo"], str(m["n"]), bs(m["ingresos"]), bs(m["egresos"]), bs(m["neto"])] for m in cuadre["por_metodo"]] or [["—", "0", bs(0), bs(0), bs(0)]],
                       [3, 1.5, 2, 2, 2], (1, 2, 3, 4), ["TOTAL", str(cuadre["n_mov"]), bs(cuadre["ingresos"]), bs(cuadre["egresos"]), bs(cuadre["resultado"])]))
    p = cuadre["pendientes"]
    story.append(Paragraph(
        f"Según el Estado (no se descuenta de los totales): egresos con devolución pendiente {bs(p['devolucion_pendiente'])}, "
        f"con devolución realizada {bs(p['devolucion_realizada'])}; ingresos con entrega pendiente {bs(p['entrega_pendiente'])}, "
        f"con entrega realizada {bs(p['entrega_realizada'])}.", st_nota))

    for titulo, clave, total in (("Egresos por categoría (pagos + compras)", "categorias_egreso", cuadre["egresos"]),
                                 ("Ingresos por concepto (ventas)", "categorias_ingreso", cuadre["ingresos"])):
        if cuadre[clave]:
            story.append(CondPageBreak(35 * mm))
            story.append(Paragraph(titulo, st_sec))
            story.append(tabla(["Categoría", "Movimientos", "Total"],
                               [[c["categoria"], str(c["n"]), bs(c["total"])] for c in cuadre[clave]],
                               [5, 1.5, 2], (1, 2), ["TOTAL", "", bs(total)]))

    # ---- Libro diario ----
    story.append(CondPageBreak(40 * mm))
    story.append(Paragraph("Libro diario de movimientos", st_sec))
    if cuadre["libro"]:
        filas = [[m["fecha_txt"], f"{m['origen']}", m["categoria"], m["detalle"], m["persona"], m["metodo_original"] or "—",
                  rp._estado_txt(m["estado"]), m["devolucion_a"], bs(m["ingreso"]) if m["ingreso"] is not None else "",
                  bs(m["egreso"]) if m["egreso"] is not None else "", bs(m["acumulado"])] for m in cuadre["libro"]]
        story.append(tabla(["Fecha", "Registro", "Categoría / Concepto", "Descripción", "Persona / Empresa", "Método", "Estado",
                            "Devolución a", "Ingreso", "Egreso", "Acumulado"], filas,
                           [1.1, 0.9, 1.6, 2.4, 1.6, 1.2, 1.6, 1.3, 1.3, 1.3, 1.4],
                           (8, 9, 10), ["TOTAL", "", "", "", "", "", "", "", bs(cuadre["ingresos"]), bs(cuadre["egresos"]), bs(cuadre["resultado"])]))
    else:
        story.append(Paragraph("Sin movimientos en este periodo.", st_nota))

    story += firmas_pdf(ancho)

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
