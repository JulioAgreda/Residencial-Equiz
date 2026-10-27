"""
Reportes de la conciliación bancaria en Excel y PDF. Reciben lo que devuelve conciliacion.conciliar() (más el resultado de la
lectura del extracto, para el control de saldos) y solo les dan formato.
"""
import io
from datetime import date

import openpyxl
from reportlab.lib.units import mm
from reportlab.platypus import CondPageBreak, Paragraph, Spacer

import conciliacion as co
import cuadre as cd
import reportes as rp


def _n(valor):
    return f"{valor:,.2f}"


def describir_sistema(s):
    quien = s.get("apartamento") or s.get("persona") or ""
    partes = [s["origen"]] + ([quien] if quien else []) + ([s["detalle"]] if s.get("detalle") and s["detalle"] != s["origen"] else [])
    return " · ".join(partes)


def describir_banco(b):
    partes = [b.get("contraparte") or b.get("transaccion") or ""] + ([b["detalle"]] if b.get("detalle") else [])
    return " · ".join(p for p in partes if p)


def estado_extracto(r):
    """{n del movimiento del banco: (estado, explicación)} para todo el extracto."""
    estados = {}
    for p in r["pares"]:
        explic = f"{p['union']} · {p['confianza']} · " + "; ".join(describir_sistema(s) + f" (Bs {_n(s['monto'])})" for s in p["sistema"])
        for b in p["banco"]:
            estados[b["n"]] = ("Conciliado", explic)
    for x in r["banco_sin_registro"]:
        estados[x["linea"]["n"]] = ("Sin registrar en el sistema", x["causas"][0])
    for b in r["excluidos"]:
        estados[b["n"]] = ("Excluido", b["motivo_exclusion"])
    return estados


# ============================================================
# Filas para las tablas (las usan la pantalla, el Excel y el PDF)
# ============================================================

def filas_conciliados(r):
    filas = []
    for p in r["pares"]:
        bancos, sistemas = p["banco"], p["sistema"]
        filas.append({
            "fecha_banco": ", ".join(sorted({co._fmt_fecha(b["fecha"]) for b in bancos})),
            "banco": " + ".join(describir_banco(b) for b in bancos), "monto_banco": round(sum(b["monto"] for b in bancos), 2),
            "fecha_sistema": ", ".join(sorted({co._fmt_fecha(s["fecha"]) for s in sistemas})),
            "sistema": " + ".join(describir_sistema(s) for s in sistemas),
            "monto_sistema": round(sum(s["monto"] if s["tipo"] == co.INGRESO else -s["monto"] for s in sistemas), 2),   # con signo, como en el banco
            "union": p["union"], "confianza": p["confianza"], "dias": p["dias"], "razones": "; ".join(p["razones"])})
    return filas


def filas_banco_sin_registro(r):
    return [{"fecha": co._fmt_fecha(x["linea"]["fecha"]), "hora": x["linea"].get("hora") or "", "nro": x["linea"].get("nro") or "",
             "transaccion": x["linea"].get("transaccion") or "", "contraparte": x["linea"].get("contraparte") or "", "detalle": x["linea"].get("detalle") or "",
             "monto": x["linea"]["monto"], "causa": x["causas"][0], "otras": " | ".join(x["causas"][1:]), "sugerencia": x["sugerencia"]}
            for x in r["banco_sin_registro"]]


def filas_sistema_sin_banco(r):
    return [{"fecha": co._fmt_fecha(x["mov"]["fecha"]), "origen": x["mov"]["origen"], "apartamento": x["mov"].get("apartamento") or "",
             "persona": x["mov"].get("persona") or "", "detalle": x["mov"].get("detalle") or "", "metodo": x["mov"].get("metodo_original") or x["mov"]["metodo"],
             "monto": x["mov"]["monto"] if x["mov"]["tipo"] == co.INGRESO else -x["mov"]["monto"], "causa": x["causas"][0],
             "otras": " | ".join(x["causas"][1:])} for x in r["sistema_sin_banco"]]


def filas_efectivo(r):
    return [{"fecha": co._fmt_fecha(s["fecha"]), "origen": s["origen"], "apartamento": s.get("apartamento") or "", "persona": s.get("persona") or "",
             "detalle": s.get("detalle") or "", "monto": s["monto"] if s["tipo"] == co.INGRESO else -s["monto"]} for s in r["efectivo"]]


def filas_extracto(r):
    est = estado_extracto(r)
    return [{"n": b["n"], "fecha": co._fmt_fecha(b["fecha"]), "hora": b.get("hora") or "", "nro": b.get("nro") or "", "transaccion": b.get("transaccion") or "",
             "contraparte": b.get("contraparte") or "", "detalle": b.get("detalle") or "", "monto": b["monto"], "saldo": b.get("saldo"),
             "estado": est[b["n"]][0], "explicacion": est[b["n"]][1]} for b in r["banco"]]


COL_CONCILIADOS = [("fecha_banco", "Fecha banco", "txt"), ("banco", "Movimiento del banco", "txt"), ("monto_banco", "Monto banco", "bs"),
                   ("fecha_sistema", "Fecha registro", "txt"), ("sistema", "Registro(s) del sistema", "txt"), ("monto_sistema", "Monto sistema", "bs"),
                   ("union", "Tipo", "txt"), ("confianza", "Confianza", "txt"), ("dias", "Días de diferencia", "num"), ("razones", "Por qué coinciden", "txt")]
COL_BANCO_SIN = [("fecha", "Fecha", "txt"), ("hora", "Hora", "txt"), ("nro", "N° transacción", "txt"), ("transaccion", "Transacción", "txt"),
                 ("contraparte", "Nota (a quién / de quién)", "txt"), ("detalle", "Detalle", "txt"), ("monto", "Monto", "bs"),
                 ("causa", "Causa probable", "txt"), ("otras", "Otras pistas", "txt"), ("sugerencia", "Qué hacer", "txt")]
COL_SISTEMA_SIN = [("fecha", "Fecha", "txt"), ("origen", "Registro", "txt"), ("apartamento", "Apartamento", "txt"), ("persona", "Persona", "txt"),
                   ("detalle", "Detalle", "txt"), ("metodo", "Método", "txt"), ("monto", "Monto", "bs"), ("causa", "Causa probable", "txt"),
                   ("otras", "Otras pistas", "txt")]
COL_EFECTIVO = [("fecha", "Fecha", "txt"), ("origen", "Registro", "txt"), ("apartamento", "Apartamento", "txt"), ("persona", "Persona", "txt"),
                ("detalle", "Detalle", "txt"), ("monto", "Monto", "bs")]
COL_EXTRACTO = [("n", "N°", "num"), ("fecha", "Fecha", "txt"), ("hora", "Hora", "txt"), ("nro", "N° transacción", "txt"), ("transaccion", "Transacción", "txt"),
                ("contraparte", "Nota (a quién / de quién)", "txt"), ("detalle", "Detalle", "txt"), ("monto", "Monto", "bs"), ("saldo", "Saldo", "bs"),
                ("estado", "Estado", "txt"), ("explicacion", "Explicación", "txt")]


def lineas_resumen(r, extracto=None):
    """[(etiqueta, valor, tipo)] para el resumen del Excel y del PDF. tipo: 'titulo', 'bs', 'num', 'txt'."""
    rs = r["resumen"]
    sub = f"Del {co._fmt_fecha(r['desde'])} al {co._fmt_fecha(r['hasta'])} · diferencia de fechas aceptada: {r['dias']} día(s) · " + \
        ("incluye los registros en efectivo" if r["incluir_efectivo"] else "sin los registros en efectivo")
    lineas = [("Periodo del extracto", sub, "txt"), ("EXTRACTO DEL BANCO", "", "titulo"),
              ("Movimientos del extracto", rs["banco"]["n"], "num"), ("Créditos (entra dinero)", rs["banco"]["creditos"], "bs"),
              ("Débitos (sale dinero)", rs["banco"]["debitos"], "bs")]
    if extracto and extracto.get("saldo_final") is not None:
        lineas += [("Saldo inicial", extracto["saldo_inicial"], "bs"), ("Saldo final", extracto["saldo_final"], "bs")]
        lineas.append(("Control del saldo del extracto", "Correcto" if not extracto["problemas_saldo"] else f"{len(extracto['problemas_saldo'])} diferencia(s): revisar", "txt"))
    lineas += [("CONCILIADO (coincide con el sistema)", "", "titulo"), ("Movimientos del banco conciliados", rs["conciliado"]["n"], "num"),
               ("Créditos conciliados", rs["conciliado"]["creditos"], "bs"), ("Débitos conciliados", rs["conciliado"]["debitos"], "bs"),
               ("Coincidencias con confianza Alta / Media / Revisar", f"{rs['por_confianza']['Alta']} / {rs['por_confianza']['Media']} / {rs['por_confianza']['Revisar']}", "txt"),
               ("EN EL BANCO, SIN REGISTRAR EN EL SISTEMA", "", "titulo"), ("Movimientos sin registrar", rs["sin_registro"]["n"], "num"),
               ("Créditos sin registrar", rs["sin_registro"]["creditos"], "bs"), ("Débitos sin registrar", rs["sin_registro"]["debitos"], "bs")]
    if rs["excluidos"]["n"]:
        lineas += [("Movimientos excluidos a propósito", rs["excluidos"]["n"], "num"), ("Créditos excluidos", rs["excluidos"]["creditos"], "bs"),
                   ("Débitos excluidos", rs["excluidos"]["debitos"], "bs")]
    lineas += [("EN EL SISTEMA, SIN MOVIMIENTO EN EL BANCO", "", "titulo"), ("Registros comparados en el periodo", rs["sistema"]["n"], "num"),
               ("Registros sin movimiento en el banco", rs["sistema_sin_banco"]["n"], "num"), ("Ingresos sin movimiento en el banco", rs["sistema_sin_banco"]["ingresos"], "bs"),
               ("Egresos sin movimiento en el banco", rs["sistema_sin_banco"]["egresos"], "bs")]
    if rs["efectivo_no_comparado"]["n"]:
        lineas += [("REGISTROS EN EFECTIVO (no se compararon con el banco)", "", "titulo"), ("Registros en efectivo", rs["efectivo_no_comparado"]["n"], "num"),
                   ("Ingresos en efectivo", rs["efectivo_no_comparado"]["ingresos"], "bs"), ("Egresos en efectivo", rs["efectivo_no_comparado"]["egresos"], "bs")]
    return lineas


# ============================================================
# EXCEL
# ============================================================

def generar_excel(r, extracto=None, emitido_por=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    from openpyxl.styles import Font, PatternFill
    f_tit = Font(name="Calibri", bold=True, size=14)
    f_sec = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    relleno = PatternFill("solid", fgColor=rp._COLOR_HEADER)
    rp._poner(ws, 1, 1, "Conciliación bancaria", f_tit)
    if emitido_por:
        rp._poner(ws, 2, 1, f"Emitido por: {emitido_por}", Font(name="Calibri", size=10))
    fila = 4
    for etiqueta, valor, tipo in lineas_resumen(r, extracto):
        if tipo == "titulo":
            c = rp._poner(ws, fila, 1, etiqueta, f_sec)
            c.fill, ws.cell(fila, 2).fill = relleno, relleno
        else:
            rp._poner(ws, fila, 1, etiqueta, Font(name="Calibri", size=10))
            rp._poner(ws, fila, 2, valor, Font(name="Calibri", size=10, bold=True), cd._FMT if tipo == "bs" else None)
        fila += 1
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 60
    sub = f"Del {co._fmt_fecha(r['desde'])} al {co._fmt_fecha(r['hasta'])}"
    cd._hoja(wb, "Conciliados", "Movimientos conciliados (banco = sistema)", sub, COL_CONCILIADOS, filas_conciliados(r))
    banco_sin = filas_banco_sin_registro(r)
    cd._hoja(wb, "Banco sin registrar", "En el banco, sin registrar en el sistema", sub, COL_BANCO_SIN, banco_sin,
             {"monto": round(sum(f["monto"] for f in banco_sin), 2)})
    sist_sin = filas_sistema_sin_banco(r)
    cd._hoja(wb, "Sistema sin banco", "Registrados en el sistema, sin movimiento en el banco", sub, COL_SISTEMA_SIN, sist_sin,
             {"monto": round(sum(f["monto"] for f in sist_sin), 2)})
    if r["efectivo"]:
        cd._hoja(wb, "Efectivo", "Registros en efectivo (no se compararon con el banco)", sub, COL_EFECTIVO, filas_efectivo(r))
    cd._hoja(wb, "Extracto", "Extracto del banco con el estado de cada movimiento", sub, COL_EXTRACTO, filas_extracto(r))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ============================================================
# PDF
# ============================================================

def generar_pdf(r, extracto=None, emitido_por=None):
    bs, pdf = rp._fmt_bs, rp._pdf_texto
    buffer = io.BytesIO()
    doc, ancho = cd.documento_pdf(buffer, "Conciliación bancaria")
    est = cd.estilos_pdf()

    def tabla(*a, **k):
        return cd.tabla_pdf(*a, ancho=ancho, est=est, **k)

    rs = r["resumen"]
    story = cd.cabecera_pdf("Conciliación bancaria", r["desde"], r["hasta"], emitido_por, ancho, est)
    story.append(Paragraph(f"Diferencia de fechas aceptada: <b>{r['dias']} día(s)</b> &nbsp;|&nbsp; "
                           + ("Incluye los registros en efectivo" if r["incluir_efectivo"] else "No incluye los registros en efectivo (no pasan por el banco)"), est.chico))
    story.append(Spacer(1, 4))
    filas_res = [["Extracto del banco", str(rs["banco"]["n"]), bs(rs["banco"]["creditos"]), bs(rs["banco"]["debitos"])],
                 ["Conciliado con el sistema", str(rs["conciliado"]["n"]), bs(rs["conciliado"]["creditos"]), bs(rs["conciliado"]["debitos"])],
                 ["En el banco, sin registrar en el sistema", str(rs["sin_registro"]["n"]), bs(rs["sin_registro"]["creditos"]), bs(rs["sin_registro"]["debitos"])]]
    if rs["excluidos"]["n"]:
        filas_res.append(["Excluidos a propósito", str(rs["excluidos"]["n"]), bs(rs["excluidos"]["creditos"]), bs(rs["excluidos"]["debitos"])])
    story.append(tabla(["Movimientos del banco", "Cantidad", "Créditos (entra)", "Débitos (sale)"], filas_res, [4, 1, 1.6, 1.6], (1, 2, 3), ancho_tabla=ancho * 0.6))
    story.append(Spacer(1, 3))
    filas_sis = [["Registrados en el sistema (comparados)", str(rs["sistema"]["n"]), bs(rs["sistema"]["ingresos"]), bs(rs["sistema"]["egresos"])],
                 ["Registrados, sin movimiento en el banco", str(rs["sistema_sin_banco"]["n"]), bs(rs["sistema_sin_banco"]["ingresos"]), bs(rs["sistema_sin_banco"]["egresos"])]]
    if rs["efectivo_no_comparado"]["n"]:
        filas_sis.append(["En efectivo (no se compararon)", str(rs["efectivo_no_comparado"]["n"]), bs(rs["efectivo_no_comparado"]["ingresos"]), bs(rs["efectivo_no_comparado"]["egresos"])])
    story.append(tabla(["Registros del sistema", "Cantidad", "Ingresos", "Egresos"], filas_sis, [4, 1, 1.6, 1.6], (1, 2, 3), ancho_tabla=ancho * 0.6))
    if extracto and extracto.get("saldo_final") is not None:
        control = "correcto" if not extracto["problemas_saldo"] else f"{len(extracto['problemas_saldo'])} diferencia(s): revisar"
        story.append(Paragraph(f"Saldo del extracto: inicial Bs {_n(extracto['saldo_inicial'])} → final Bs {_n(extracto['saldo_final'])} (control del saldo: {control}).", est.nota))

    banco_sin = filas_banco_sin_registro(r)
    story.append(CondPageBreak(40 * mm))
    story.append(Paragraph(f"En el banco, sin registrar en el sistema ({len(banco_sin)})", est.sec))
    if banco_sin:
        story.append(tabla(["Fecha", "Nota (a quién / de quién)", "Detalle", "Monto", "Causa probable", "Qué hacer"],
                           [[f["fecha"], f["contraparte"], f["detalle"], bs(f["monto"]), f["causa"], f["sugerencia"]] for f in banco_sin],
                           [1.0, 2.2, 2.4, 1.2, 5.0, 2.4], (3,)))
    else:
        story.append(Paragraph("Todos los movimientos del banco están registrados en el sistema.", est.nota))
    sist_sin = filas_sistema_sin_banco(r)
    story.append(CondPageBreak(40 * mm))
    story.append(Paragraph(f"Registrados en el sistema, sin movimiento en el banco ({len(sist_sin)})", est.sec))
    if sist_sin:
        story.append(tabla(["Fecha", "Registro", "Apartamento / persona", "Detalle", "Método", "Monto", "Causa probable"],
                           [[f["fecha"], f["origen"], f["apartamento"] or f["persona"], f["detalle"], f["metodo"], bs(f["monto"]), f["causa"]] for f in sist_sin],
                           [1.0, 1.1, 2.2, 2.2, 1.1, 1.2, 5.0], (5,)))
    else:
        story.append(Paragraph("Todos los registros del sistema aparecen en el banco.", est.nota))
    conc = filas_conciliados(r)
    story.append(CondPageBreak(40 * mm))
    story.append(Paragraph(f"Conciliados ({len(conc)} coincidencias)", est.sec))
    if conc:
        story.append(tabla(["Fecha banco", "Movimiento del banco", "Monto", "Registro(s) del sistema", "Tipo", "Confianza"],
                           [[f["fecha_banco"], f["banco"], bs(f["monto_banco"]), f["sistema"], f["union"], f["confianza"]] for f in conc],
                           [1.1, 3.4, 1.2, 4.4, 1.8, 1.0], (2,)))
        story.append(Paragraph("«Confianza»: Alta = el texto confirma la coincidencia; Media = coincide el monto y la fecha pero el texto no lo confirma o hay varias "
                               "posibilidades; Revisar = posible coincidencia que conviene verificar.", est.nota))
    else:
        story.append(Paragraph("No hubo coincidencias.", est.nota))
    story += cd.firmas_pdf(ancho, "Elaborado por", "Revisado por (Administración)")
    doc.build(story)
    return buffer.getvalue()
